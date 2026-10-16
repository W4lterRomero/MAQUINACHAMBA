"""
Generador local de documentos de CV en formatos PDF y DOCX (REQ-013, D9, D12).
Principios clave:
1. Local-first y determinista: generación puramente offline sin llamadas a servicios externos.
2. Fidelidad a hechos verificados: plasma exactamente las habilidades y experiencias aprobadas,
   sin inventar datos, años ni empleos ficticios.
3. Preservación tipográfica: manejo riguroso de caracteres UTF-8 (acentos, ñ, signos de puntuación).
4. Estructura visual legible: paginación controlada, espaciado proporcional y jerarquía clara.
"""

from datetime import date, datetime
import io
from typing import Any, Dict, List, Optional
import xml.sax.saxutils

# ReportLab para PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# python-docx para DOCX
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def _esc(text: Optional[str]) -> str:
    """Escapa entidades XML/HTML para párrafos Platypus de ReportLab preservando caracteres UTF-8."""
    if text is None:
        return ""
    return xml.sax.saxutils.escape(str(text))


class NumberedCanvas(canvas.Canvas):
    """Canvas de dos pasadas para calcular y numerar páginas dinámicamente."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#718096"))
        page_text = f"Página {self._pageNumber} de {page_count}"
        self.drawRightString(letter[0] - 36, 24, page_text)
        footer_note = "Documento verificado localmente — Máquina de Chamba"
        self.drawString(36, 24, footer_note)
        self.restoreState()


class CVDocumentGenerator:
    """Generador local y seguro de documentos curriculares en PDF y DOCX."""

    @classmethod
    def generate_pdf(cls, cv_data: Dict[str, Any]) -> bytes:
        """
        Genera un archivo PDF estructurado y estilizado a partir de los datos verificados del CV.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=45,
        )

        base_styles = getSampleStyleSheet()

        # Estilos tipográficos personalizados
        name_style = ParagraphStyle(
            "CVName",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#1A365D"),
        )
        subtitle_style = ParagraphStyle(
            "CVSubtitle",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#2B6CB0"),
        )
        contact_style = ParagraphStyle(
            "CVContact",
            parent=base_styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#4A5568"),
        )
        section_heading = ParagraphStyle(
            "CVSectionHeading",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#2D3748"),
            spaceBefore=8,
            spaceAfter=4,
        )
        body_style = ParagraphStyle(
            "CVBody",
            parent=base_styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor("#2D3748"),
        )
        item_title_style = ParagraphStyle(
            "CVItemTitle",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#1A202C"),
        )
        item_meta_style = ParagraphStyle(
            "CVItemMeta",
            parent=base_styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#718096"),
        )

        story = []

        # 1. Cabecera (Nombre, especialización y datos de contacto)
        full_name = cv_data.get("full_name") or cv_data.get("name") or "Candidato"
        story.append(Paragraph(_esc(full_name), name_style))

        specialization = cv_data.get("specialization") or cv_data.get("title")
        if specialization:
            story.append(Spacer(1, 2))
            story.append(Paragraph(_esc(specialization), subtitle_style))

        contact_parts = []
        if cv_data.get("email"):
            contact_parts.append(cv_data["email"])
        if cv_data.get("phone"):
            contact_parts.append(cv_data["phone"])
        location = cv_data.get("location") or cv_data.get("location_city")
        if location:
            contact_parts.append(location)

        if contact_parts:
            story.append(Spacer(1, 3))
            contact_line = "  •  ".join(_esc(p) for p in contact_parts)
            story.append(Paragraph(contact_line, contact_style))

        story.append(Spacer(1, 6))
        story.append(
            HRFlowable(
                width="100%",
                thickness=1.5,
                color=colors.HexColor("#3182CE"),
                spaceBefore=0,
                spaceAfter=8,
            )
        )

        # 2. Resumen Profesional
        summary = cv_data.get("summary")
        if summary and summary.strip():
            story.append(Paragraph("Perfil Profesional", section_heading))
            story.append(Paragraph(_esc(summary.strip()), body_style))
            story.append(Spacer(1, 6))

        # 3. Competencias Técnicas Verificadas
        raw_skills = cv_data.get("skills") or cv_data.get("verified_skills") or []
        if raw_skills:
            story.append(Paragraph("Competencias Técnicas Verificadas", section_heading))

            formatted_skills = []
            for sk in raw_skills:
                if isinstance(sk, dict):
                    sk_name = sk.get("name", "")
                    prof = sk.get("proficiency")
                    years = sk.get("years_experience")
                    details = []
                    if prof:
                        details.append(prof.capitalize())
                    if years:
                        details.append(f"{years} a.")
                    if details:
                        formatted_skills.append(f"<b>{_esc(sk_name)}</b> ({_esc(', '.join(details))})")
                    else:
                        formatted_skills.append(f"<b>{_esc(sk_name)}</b>")
                else:
                    formatted_skills.append(f"<b>{_esc(str(sk))}</b>")

            # Formatear habilidades en una lista limpia de 2 columnas
            col1 = []
            col2 = []
            for i, item in enumerate(formatted_skills):
                bullet_item = Paragraph(f"• {item}", body_style)
                if i % 2 == 0:
                    col1.append(bullet_item)
                else:
                    col2.append(bullet_item)

            max_rows = max(len(col1), len(col2))
            table_data = []
            for r in range(max_rows):
                c1 = col1[r] if r < len(col1) else Paragraph("", body_style)
                c2 = col2[r] if r < len(col2) else Paragraph("", body_style)
                table_data.append([c1, c2])

            skills_table = Table(
                table_data,
                colWidths=[270, 270],
                spaceBefore=2,
                spaceAfter=4,
            )
            skills_table.setStyle(
                TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 1),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                ])
            )
            story.append(skills_table)
            story.append(Spacer(1, 6))

        # 4. Trayectoria Profesional Verificada
        experiences = cv_data.get("experiences") or cv_data.get("verified_experiences") or []
        if experiences:
            story.append(Paragraph("Trayectoria Profesional", section_heading))

            for exp in experiences:
                exp_flowables = []
                title = exp.get("title") or "Puesto profesional"
                company = exp.get("company") or "Empresa"
                header_text = f"<b>{_esc(title)}</b> — <i>{_esc(company)}</i>"
                exp_flowables.append(Paragraph(header_text, item_title_style))

                start = exp.get("start_date") or ""
                end = exp.get("end_date") or "Presente"
                date_meta = f"{_esc(str(start))} – {_esc(str(end))}"
                exp_flowables.append(Paragraph(date_meta, item_meta_style))

                desc = exp.get("description")
                if desc and desc.strip():
                    exp_flowables.append(Spacer(1, 2))
                    exp_flowables.append(Paragraph(_esc(desc.strip()), body_style))

                skills_used = exp.get("skills_used")
                if isinstance(skills_used, list) and skills_used:
                    skills_line = "Herramientas aplicadas: " + ", ".join(_esc(str(s)) for s in skills_used)
                    exp_flowables.append(Spacer(1, 1))
                    exp_flowables.append(Paragraph(f"<i>{skills_line}</i>", item_meta_style))

                exp_flowables.append(Spacer(1, 5))
                story.append(KeepTogether(exp_flowables))

        # 5. Formación Académica
        education = cv_data.get("education") or []
        if education:
            story.append(Paragraph("Formación Académica", section_heading))
            for edu in education:
                edu_flowables = []
                degree = edu.get("degree") or "Título"
                inst = edu.get("institution") or "Institución"
                edu_flowables.append(
                    Paragraph(f"<b>{_esc(degree)}</b> — {_esc(inst)}", item_title_style)
                )
                start = edu.get("start_date") or ""
                end = edu.get("end_date") or ""
                if start or end:
                    edu_flowables.append(
                        Paragraph(f"{_esc(str(start))} – {_esc(str(end))}", item_meta_style)
                    )
                edu_flowables.append(Spacer(1, 4))
                story.append(KeepTogether(edu_flowables))

        doc.build(story, canvasmaker=NumberedCanvas)
        return buffer.getvalue()

    @classmethod
    def generate_docx(cls, cv_data: Dict[str, Any]) -> bytes:
        """
        Genera un archivo DOCX estructurado a partir de los datos verificados del CV.
        """
        doc = docx.Document()

        # Configurar márgenes de 0.6 pulgadas
        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        # 1. Cabecera (Nombre)
        full_name = cv_data.get("full_name") or cv_data.get("name") or "Candidato"
        h1 = doc.add_heading(level=1)
        r1 = h1.add_run(full_name)
        r1.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        r1.font.size = Pt(20)
        r1.bold = True
        h1.paragraph_format.space_after = Pt(2)

        # Especialización
        specialization = cv_data.get("specialization") or cv_data.get("title")
        if specialization:
            p_spec = doc.add_paragraph()
            r_spec = p_spec.add_run(specialization)
            r_spec.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            r_spec.font.size = Pt(11)
            r_spec.bold = True
            p_spec.paragraph_format.space_after = Pt(2)

        # Contacto
        contact_parts = []
        if cv_data.get("email"):
            contact_parts.append(cv_data["email"])
        if cv_data.get("phone"):
            contact_parts.append(cv_data["phone"])
        location = cv_data.get("location") or cv_data.get("location_city")
        if location:
            contact_parts.append(location)

        if contact_parts:
            p_contact = doc.add_paragraph()
            r_contact = p_contact.add_run("  •  ".join(contact_parts))
            r_contact.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
            r_contact.font.size = Pt(9.5)
            p_contact.paragraph_format.space_after = Pt(8)

        # 2. Resumen Profesional
        summary = cv_data.get("summary")
        if summary and summary.strip():
            h_sum = doc.add_heading("Perfil Profesional", level=2)
            h_sum.paragraph_format.space_before = Pt(8)
            h_sum.paragraph_format.space_after = Pt(3)
            p_sum = doc.add_paragraph(summary.strip())
            p_sum.paragraph_format.space_after = Pt(6)

        # 3. Competencias Técnicas Verificadas
        raw_skills = cv_data.get("skills") or cv_data.get("verified_skills") or []
        if raw_skills:
            h_sk = doc.add_heading("Competencias Técnicas Verificadas", level=2)
            h_sk.paragraph_format.space_before = Pt(8)
            h_sk.paragraph_format.space_after = Pt(3)

            for sk in raw_skills:
                p_sk = doc.add_paragraph(style="List Bullet")
                if isinstance(sk, dict):
                    sk_name = sk.get("name", "")
                    prof = sk.get("proficiency")
                    years = sk.get("years_experience")
                    r_name = p_sk.add_run(sk_name)
                    r_name.bold = True
                    details = []
                    if prof:
                        details.append(prof.capitalize())
                    if years:
                        details.append(f"{years} a.")
                    if details:
                        p_sk.add_run(f" ({', '.join(details)})")
                else:
                    r_name = p_sk.add_run(str(sk))
                    r_name.bold = True
                p_sk.paragraph_format.space_after = Pt(2)

        # 4. Trayectoria Profesional
        experiences = cv_data.get("experiences") or cv_data.get("verified_experiences") or []
        if experiences:
            h_exp = doc.add_heading("Trayectoria Profesional", level=2)
            h_exp.paragraph_format.space_before = Pt(8)
            h_exp.paragraph_format.space_after = Pt(3)

            for exp in experiences:
                title = exp.get("title") or "Puesto profesional"
                company = exp.get("company") or "Empresa"
                p_item = doc.add_paragraph()
                r_title = p_item.add_run(title)
                r_title.bold = True
                p_item.add_run(" — ")
                r_comp = p_item.add_run(company)
                r_comp.italic = True
                p_item.paragraph_format.space_after = Pt(1)

                start = exp.get("start_date") or ""
                end = exp.get("end_date") or "Presente"
                p_meta = doc.add_paragraph()
                r_meta = p_meta.add_run(f"{start} – {end}")
                r_meta.font.size = Pt(8.5)
                r_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
                p_meta.paragraph_format.space_after = Pt(3)

                desc = exp.get("description")
                if desc and desc.strip():
                    p_desc = doc.add_paragraph(desc.strip())
                    p_desc.paragraph_format.space_after = Pt(4)

        # 5. Formación Académica
        education = cv_data.get("education") or []
        if education:
            h_edu = doc.add_heading("Formación Académica", level=2)
            h_edu.paragraph_format.space_before = Pt(8)
            h_edu.paragraph_format.space_after = Pt(3)
            for edu in education:
                degree = edu.get("degree") or "Título"
                inst = edu.get("institution") or "Institución"
                p_edu = doc.add_paragraph()
                r_deg = p_edu.add_run(degree)
                r_deg.bold = True
                p_edu.add_run(" — ")
                p_edu.add_run(inst)
                p_edu.paragraph_format.space_after = Pt(2)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
